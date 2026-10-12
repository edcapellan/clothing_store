"""Data layer.

Products come from a SPREADSHEET (the single source of truth):

  1. If the PRODUCTS_SHEET_URL environment variable is set, products are
     read from that Google Sheet (published/shared as CSV) and refreshed
     every few minutes -- edit the sheet, the site updates, no redeploy.
  2. Otherwise products are read from products.xlsx (or products.csv) that
     lives next to this file. Edit it, commit, push, done.

If the Google Sheet can't be reached, the last good copy is kept and the
local file is used as a fallback, so a Google outage never empties the shop.

Orders are appended to orders.json (see the note on DATA_DIR below).

Everything here is plain Python 3.9+ -- no pandas/numpy -- so it installs
fast on Render.
"""

import csv
import io
import json
import logging
import os
import re
import threading
import time
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

log = logging.getLogger("noemi.data")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
IMAGES_DIR = os.path.join(BASE_DIR, "static", "images")

PRODUCTS_XLSX = os.path.join(BASE_DIR, "products.xlsx")
PRODUCTS_CSV = os.path.join(BASE_DIR, "products.csv")

SHEET_URL = os.getenv("PRODUCTS_SHEET_URL", "").strip()
SHEET_CACHE_SECONDS = int(os.getenv("PRODUCTS_CACHE_SECONDS", "300"))

# Orders live in orders.json. On hosts with an ephemeral disk (Render's free
# tier) that file is wiped on every deploy/restart. Point DATA_DIR at a
# persistent disk mount (e.g. /var/data) to keep it. Stripe's dashboard is
# always the permanent record of payments either way.
DATA_DIR = os.getenv("DATA_DIR", BASE_DIR)
ORDERS_FILE = os.path.join(DATA_DIR, "orders.json")

# Keep in sync with CATEGORY_META in category_routes.py.
VALID_CATEGORIES = ("nails", "accessories", "clothing", "jewelry")
CATEGORY_DISPLAY = {
    "nails": "Nails",
    "accessories": "Accessories",
    "clothing": "Clothing",
    "jewelry": "Jewelry",
}
_CATEGORY_ALIASES = {
    "nail": "nails",
    "accessory": "accessories",
    "jewellery": "jewelry",
    "jewelery": "jewelry",
    "jewlery": "jewelry",
}

# Spreadsheet column names are matched loosely ("In Stock", "in_stock" and
# "IN STOCK" are all the same column). A few friendly alternatives too.
_HEADER_ALIASES = {
    "product": "name",
    "product_name": "name",
    "title": "name",
    "sub_category": "subcategory",
    "type": "subcategory",
    "images": "image",
    "photo": "image",
    "image_file": "image",
    "image_url": "image",
    "desc": "description",
    "stock": "in_stock",
    "instock": "in_stock",
    "available": "in_stock",
    "color": "placeholder_color",
    "colour": "placeholder_color",
    "show": "visible",
    "show_on_site": "visible",
    "active": "visible",
    "published": "visible",
}

_TRUE_WORDS = {"yes", "y", "true", "t", "1", "x", "\u2713", "\u2714", "in stock", "available"}
_FALSE_WORDS = {"no", "n", "false", "f", "0", "-", "sold out", "soldout", "out of stock", "out"}

_IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".webp", ".gif")

_lock = threading.Lock()
_cache = {
    "products": [],
    "warnings": [],
    "stamp": None,
    "source": "not loaded yet",
    "loaded_at": 0.0,
}


# ---------------------------------------------------------------------------
# Small parsing helpers
# ---------------------------------------------------------------------------
def _norm_header(value):
    key = re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower()).strip("_")
    return _HEADER_ALIASES.get(key, key)


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _slugify(text):
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return text or "item"


def _parse_price(value):
    """Return a Decimal with 2 places. Accepts 24.99, '$24.99', '1,299.50', '24,99'."""
    if value is None or str(value).strip() == "":
        raise ValueError("price is empty")
    if isinstance(value, bool):
        raise ValueError("price is not a number")
    if isinstance(value, (int, float)):
        number = Decimal(str(value))
    else:
        s = str(value).strip().replace("$", "").replace("USD", "").replace(" ", "")
        if "," in s and "." in s:
            s = s.replace(",", "")
        elif "," in s:
            # "24,99" is a decimal comma; "1,299" is a thousands comma
            s = s.replace(",", ".") if re.fullmatch(r"\d+,\d{1,2}", s) else s.replace(",", "")
        try:
            number = Decimal(s)
        except InvalidOperation:
            raise ValueError("price '%s' is not a number" % value)
    number = number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if number <= 0:
        raise ValueError("price must be greater than 0")
    return number


def _parse_bool(value, default):
    if isinstance(value, bool):
        return value
    s = _text(value).lower()
    if s == "":
        return default
    if s in _TRUE_WORDS:
        return True
    if s in _FALSE_WORDS:
        return False
    raise ValueError("'%s' is not Yes/No" % value)


_COLOR_RE = re.compile(r"^#(?:[0-9a-fA-F]{3}|[0-9a-fA-F]{6})$")


def _parse_color(value):
    # Strictly validated because it is written into a style="" attribute.
    s = _text(value)
    return s if _COLOR_RE.match(s) else None


# ---------------------------------------------------------------------------
# Image resolution (case-insensitive: Windows doesn't care about capitals,
# Render's Linux does -- "Set1.JPG" would work locally and 404 in production)
# ---------------------------------------------------------------------------
def _image_index():
    index = {}
    if not os.path.isdir(IMAGES_DIR):
        return index
    for root, _dirs, files in os.walk(IMAGES_DIR):
        for name in files:
            rel = os.path.relpath(os.path.join(root, name), IMAGES_DIR).replace(os.sep, "/")
            index[rel.lower()] = rel
    return index


def _resolve_image(raw, index, notes, label):
    """Return (image_path_or_url, ok). Paths are relative to /static."""
    raw = _text(raw)
    if not raw:
        return None, False
    if re.match(r"^https?://", raw, re.I):
        return raw, True

    name = raw.replace("\\", "/")
    name = re.sub(r"^/?(static/)?(images/)?", "", name, flags=re.I)
    if ".." in name.split("/"):
        notes.append("%s: image '%s' is not allowed" % (label, raw))
        return None, False

    candidates = [name]
    if not os.path.splitext(name)[1]:
        candidates += [name + ext for ext in _IMAGE_EXTENSIONS]
    for candidate in candidates:
        actual = index.get(candidate.lower())
        if actual:
            if actual != candidate:
                notes.append("%s: image '%s' matched '%s' (fix the spelling/capitals in the "
                             "sheet)" % (label, raw, actual))
            return "images/" + actual, True

    notes.append("%s: image '%s' not found in static/images" % (label, raw))
    return None, False


# ---------------------------------------------------------------------------
# Reading the spreadsheet into raw rows: [(row_number, {column: value}), ...]
# ---------------------------------------------------------------------------
def _rows_from_iterable(rows):
    header = None
    out = []
    for number, row in enumerate(rows, start=1):
        if header is None:
            if row is None or all(_text(c) == "" for c in row):
                continue  # skip blank lines above the header
            header = [_norm_header(c) if _text(c) != "" else None for c in row]
            continue
        record = {}
        for col, value in zip(header, row):
            if col:
                record[col] = value
        out.append((number, record))
    if header is None:
        raise ValueError("the spreadsheet is empty")
    return out


def _read_xlsx(path):
    from openpyxl import load_workbook  # imported here so CSV/Sheet mode don't need it

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
        return _rows_from_iterable(ws.iter_rows(values_only=True))
    finally:
        wb.close()


def _read_csv_text(text):
    return _rows_from_iterable(csv.reader(io.StringIO(text)))


def _fetch_sheet_text():
    import requests

    resp = requests.get(SHEET_URL, timeout=10)
    resp.raise_for_status()
    text = resp.content.decode("utf-8-sig")
    if "<html" in text[:500].lower():
        raise ValueError("Google returned a web page, not CSV. Share the sheet as "
                         "'Anyone with the link: Viewer' or use File > Share > Publish to web.")
    return text


# ---------------------------------------------------------------------------
# Turning raw rows into validated products
# ---------------------------------------------------------------------------
def _build_products(raw_rows):
    index = _image_index()
    products, notes, seen = [], [], set()

    for row_number, raw in raw_rows:
        if all(_text(v) == "" for v in raw.values()):
            continue  # fully blank row

        name = _text(raw.get("name"))
        label = "Row %d" % row_number
        if not name:
            notes.append("%s: has data but no Name -- skipped" % label)
            continue
        label = "Row %d (%s)" % (row_number, name)

        try:
            visible = _parse_bool(raw.get("visible"), True)
        except ValueError as exc:
            notes.append("%s: Visible %s -- treated as Yes" % (label, exc))
            visible = True
        if not visible:
            continue

        try:
            price = _parse_price(raw.get("price"))
        except ValueError as exc:
            notes.append("%s: %s -- product skipped" % (label, exc))
            continue

        cat_raw = _text(raw.get("category")).lower()
        category = _CATEGORY_ALIASES.get(cat_raw, cat_raw)
        if category not in VALID_CATEGORIES:
            notes.append("%s: category '%s' must be one of %s -- product skipped"
                         % (label, _text(raw.get("category")), ", ".join(CATEGORY_DISPLAY.values())))
            continue

        try:
            in_stock = _parse_bool(raw.get("in_stock"), True)
        except ValueError as exc:
            notes.append("%s: In Stock %s -- treated as Yes" % (label, exc))
            in_stock = True
        try:
            featured = _parse_bool(raw.get("featured"), False)
        except ValueError as exc:
            notes.append("%s: Featured %s -- treated as No" % (label, exc))
            featured = False

        image, image_ok = _resolve_image(raw.get("image"), index, notes, label)

        slug = base = _slugify(name)
        n = 2
        while slug in seen:
            slug = "%s-%d" % (base, n)
            n += 1
        if slug != base:
            notes.append("%s: another product already has this name -- it was given the "
                         "id '%s'" % (label, slug))
        seen.add(slug)

        products.append({
            "id": len(products) + 1,
            "slug": slug,
            "name": name,
            "category": CATEGORY_DISPLAY[category],
            "subcategory": _text(raw.get("subcategory")),
            "price": float(price),
            "price_cents": int(price * 100),
            "image": image,
            "image_ok": image_ok,
            "description": _text(raw.get("description"))[:800],
            "in_stock": in_stock,
            "featured": featured,
            "placeholder_color": _parse_color(raw.get("placeholder_color")),
            "row": row_number,
        })
    return products, notes


# ---------------------------------------------------------------------------
# Public API: products
# ---------------------------------------------------------------------------
def _local_source():
    for path in (PRODUCTS_XLSX, PRODUCTS_CSV):
        if os.path.exists(path):
            return path
    return None


def _local_stamp():
    path = _local_source()
    if not path:
        return None
    stat = os.stat(path)
    return (path, stat.st_mtime, stat.st_size)


def _load_local():
    path = _local_source()
    if not path:
        raise FileNotFoundError("No products.xlsx or products.csv found next to app.py")
    if path.endswith(".xlsx"):
        return _read_xlsx(path), os.path.basename(path)
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        return _read_csv_text(fh.read()), os.path.basename(path)


def load_products(force=False):
    """Return the current product list (a fresh list each call).

    Cheap to call on every request: results are cached and only re-read when
    the local file changes or, in Google Sheet mode, every SHEET_CACHE_SECONDS.
    """
    with _lock:
        now = time.time()
        if SHEET_URL:
            fresh = (now - _cache["loaded_at"]) < SHEET_CACHE_SECONDS
            stamp = ("sheet", SHEET_URL)
        else:
            stamp = _local_stamp()
            fresh = stamp is not None and stamp == _cache["stamp"]
        if fresh and not force and _cache["stamp"] is not None:
            return list(_cache["products"])

        problems = []
        raw_rows, source = None, None

        if SHEET_URL:
            try:
                raw_rows, source = _read_csv_text(_fetch_sheet_text()), "Google Sheet"
            except Exception as exc:  # network, bad share settings, malformed CSV...
                log.warning("Could not load Google Sheet: %s", exc)
                problems.append("Could not load the Google Sheet (%s). Showing the last good "
                                "copy / local file instead." % exc)
                if _cache["products"]:
                    # keep serving what we had, retry in 30s instead of every request
                    _cache["warnings"] = problems + [w for w in _cache["warnings"]
                                                     if not w.startswith("Could not load")]
                    _cache["loaded_at"] = now - SHEET_CACHE_SECONDS + 30
                    return list(_cache["products"])

        if raw_rows is None:
            try:
                raw_rows, source = _load_local()
            except Exception as exc:
                log.error("Could not load products: %s", exc)
                problems.append("Could not load products: %s" % exc)
                _cache.update(warnings=problems, source="error", loaded_at=now, stamp=stamp)
                return list(_cache["products"])

        products, notes = _build_products(raw_rows)
        _cache.update(
            products=products,
            warnings=problems + notes,
            stamp=stamp,
            source=source,
            loaded_at=now,
        )
        for note in notes:
            log.warning("products: %s", note)
        return list(products)


def get_load_report():
    """Info for the admin page: where products came from + any row problems."""
    load_products()
    with _lock:
        return {
            "source": _cache["source"],
            "product_count": len(_cache["products"]),
            "warnings": list(_cache["warnings"]),
        }


def get_product(slug):
    for product in load_products():
        if product["slug"] == slug:
            return product
    return None


# ---------------------------------------------------------------------------
# Public API: orders
# ---------------------------------------------------------------------------
_orders_lock = threading.Lock()


def load_orders():
    if not os.path.exists(ORDERS_FILE):
        return []
    try:
        with open(ORDERS_FILE, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except ValueError:
        # Never silently overwrite a damaged file: set it aside and start fresh.
        backup = "%s.corrupt-%d" % (ORDERS_FILE, int(time.time()))
        os.replace(ORDERS_FILE, backup)
        log.error("orders.json was unreadable; moved to %s", backup)
        return []


def _write_orders(orders):
    os.makedirs(DATA_DIR, exist_ok=True)
    tmp = ORDERS_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(orders, fh, indent=2)
    os.replace(tmp, ORDERS_FILE)  # atomic: a crash can't leave half a file


def record_order(order):
    """Append an order once. Returns (order, created).

    If an order with the same payment_intent already exists (customer hit
    refresh on the success page) the existing one is returned untouched,
    so a paid order is never recorded twice.
    """
    with _orders_lock:
        orders = load_orders()
        pi = order.get("payment_intent")
        if pi:
            for existing in orders:
                if existing.get("payment_intent") == pi:
                    return existing, False
        order = dict(order, id=max([o.get("id", 0) for o in orders] or [0]) + 1)
        orders.append(order)
        _write_orders(orders)
        return order, True
